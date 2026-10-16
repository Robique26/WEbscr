"""Bot Telegram pentru monitorizare prețuri.

Pornire:  python bot.py
Produsele se adaugă direct din Telegram: trimite un link sau /add <link> [nume].
"""
import asyncio
import csv
import html
import io
import logging
import time
from logging.handlers import RotatingFileHandler
import re
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

from curl_cffi.requests import AsyncSession
from telegram import LinkPreviewOptions, Update
from telegram.constants import ParseMode
from telegram.ext import Application, CommandHandler, ContextTypes, Defaults, MessageHandler, filters

import config
import db
from scraper import Scraper

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(name)s: %(message)s",
    handlers=[
        logging.StreamHandler(),
        RotatingFileHandler(config.LOG_FILE, maxBytes=2_000_000, backupCount=3, encoding="utf-8"),
    ],
)
logging.getLogger("httpx").setLevel(logging.WARNING)
log = logging.getLogger("pricebot")

check_lock = asyncio.Lock()
started_at = time.time()

HELP = """<b>🛒 Monitor prețuri</b>

Trimite-mi pur și simplu un <b>link de produs</b> și îl urmăresc.

/add &lt;link&gt; [nume] – adaugă produs
/lista – produsele urmărite (ultimul preț)
/pret – verifică ACUM toate prețurile
/tinta &lt;id&gt; &lt;preț&gt; – alertă când ajunge sub preț (0 = șterge)
/istoric &lt;id&gt; – istoricul prețului
/selector &lt;id&gt; &lt;css&gt; – selector CSS manual (rar necesar)
/sterge &lt;id&gt; – nu mai urmări produsul
/import &lt;link CSV Google Sheets&gt; – import în masă
/status – starea botului

📎 Poți trimite și un fișier <b>.xlsx</b>/<b>.csv</b> cu coloanele <code>nume, url, selector</code>."""


# ---------------------------------------------------------------- utilitare


def fmt_price(price, currency=None):
    if price is None:
        return "—"
    return f"{price:,.2f}".replace(",", " ") + (f" {currency}" if currency else "")


def link(row_or_name, url=None):
    if url is None:
        row_or_name, url = row_or_name["name"], row_or_name["url"]
    return f'<a href="{html.escape(url)}">{html.escape(row_or_name)}</a>'


def stock_icon(in_stock):
    return {1: "", 0: " ❌ stoc epuizat"}.get(in_stock, "")


async def send_long(bot, chat_id, lines, header=""):
    """Telegram acceptă max 4096 caractere/mesaj -> împărțim pe linii."""
    chunk = header
    for line in lines:
        if len(chunk) + len(line) > 3900:
            await bot.send_message(chat_id, chunk)
            chunk = ""
        chunk += line + "\n"
    if chunk.strip():
        await bot.send_message(chat_id, chunk)


def restricted(handler):
    async def wrapper(update: Update, context: ContextTypes.DEFAULT_TYPE):
        if config.ALLOWED_CHAT_IDS and update.effective_chat.id not in config.ALLOWED_CHAT_IDS:
            log.warning("Acces refuzat pentru chat %s", update.effective_chat.id)
            return
        return await handler(update, context)

    return wrapper


_TRACKING = re.compile(r"^(utm_|x-|gclid|gbraid|wbraid|fbclid|_gl|ref|srsltid|mc_)", re.I)


def clean_url(url):
    """Scoate parametrii de tracking și #ancora -> același produs nu apare de 2 ori."""
    parts = urlsplit(url.strip())
    query = urlencode([(k, v) for k, v in parse_qsl(parts.query, keep_blank_values=True) if not _TRACKING.match(k)])
    return urlunsplit((parts.scheme, parts.netloc, parts.path, query, ""))


def parse_id(context):
    try:
        return int(context.args[0])
    except (IndexError, ValueError):
        return None


# ---------------------------------------------------------------- verificare + alerte


def build_alerts(row, res):
    old, new, cur = row["last_price"], res.price, res.currency or row["currency"]
    name = link(row)
    alerts = []
    if old and new != old:
        pct = (new - old) / old * 100
        if abs(pct) >= config.MIN_CHANGE_PCT:
            if new < old:
                stats = db.price_stats(row["id"])
                record = " 🏆 <b>minim istoric!</b>" if stats["n"] > 1 and new < stats["min"] else ""
                alerts.append(f"📉 <b>IEFTINIRE</b> {name}\n{fmt_price(old, cur)} → <b>{fmt_price(new, cur)}</b> ({pct:+.1f}%){record}")
            elif config.NOTIFY_ON_RISE:
                alerts.append(f"📈 <b>SCUMPIRE</b> {name}\n{fmt_price(old, cur)} → {fmt_price(new, cur)} ({pct:+.1f}%)")
    target = row["target_price"]
    if target and new <= target and (old is None or old > target):
        alerts.append(f"🎯 <b>PREȚ ȚINTĂ ATINS</b> {name}\n<b>{fmt_price(new, cur)}</b> (ținta ta: {fmt_price(target, cur)})")
    if row["in_stock"] == 0 and res.in_stock is True:
        alerts.append(f"✅ <b>ÎNAPOI ÎN STOC</b> {name} – {fmt_price(new, cur)}")
    elif row["in_stock"] == 1 and res.in_stock is False:
        alerts.append(f"❌ <b>STOC EPUIZAT</b> {name}")
    return alerts


async def check_product(app, row, notify=True):
    res = await app.bot_data["scraper"].scrape(row["url"], row["selector"])
    if not res.price:
        fails = db.record_failure(row["id"])
        if notify and fails == config.FAILS_BEFORE_WARNING:
            await app.bot.send_message(
                row["chat_id"],
                f"⚠️ Nu reușesc să citesc prețul pentru {link(row)} de {fails} ori ({res.error}).\n"
                f"Încearcă /selector {row['id']} &lt;css&gt; sau verifică linkul.",
            )
        return res
    alerts = build_alerts(row, res) if notify else []
    in_stock = None if res.in_stock is None else int(res.in_stock)
    db.record_success(row["id"], res.price, res.currency, in_stock, res.method)
    for text in alerts:
        await app.bot.send_message(row["chat_id"], text)
    return res


async def check_rows(app, rows, notify=True):
    async def safe(row):
        try:
            return await check_product(app, row, notify)
        except Exception:
            log.exception("Eroare la verificarea produsului %s", row["id"])
            return None

    return await asyncio.gather(*(safe(r) for r in rows))


async def check_all_job(context: ContextTypes.DEFAULT_TYPE):
    if check_lock.locked():
        return
    async with check_lock:
        rows = db.list_products()
        t0 = time.monotonic()
        results = await check_rows(context.application, rows)
        ok = sum(1 for r in results if r and r.price)
        log.info("Verificare: %d/%d prețuri în %.1fs", ok, len(rows), time.monotonic() - t0)


async def add_items(app, chat_id, items):
    """items = [{'url', 'name'|None, 'selector'|None}] -> scraping în paralel + salvare."""
    scraper = app.bot_data["scraper"]
    for item in items:
        item["url"] = clean_url(item["url"])
    results = await asyncio.gather(*(scraper.scrape(i["url"], i.get("selector")) for i in items))
    lines = []
    for item, res in zip(items, results):
        name = item.get("name") or res.title or urlsplit(item["url"]).netloc
        pid = db.add_product(chat_id, name, item["url"], item.get("selector"))
        if res.price:
            db.record_success(pid, res.price, res.currency, None if res.in_stock is None else int(res.in_stock), res.method)
            lines.append(f"✅ <code>#{pid}</code> {link(name, item['url'])}: <b>{fmt_price(res.price, res.currency)}</b>")
        else:
            lines.append(f"⚠️ <code>#{pid}</code> {link(name, item['url'])}: preț negăsit ({res.error}) – încearcă /selector {pid}")
    return lines


# ---------------------------------------------------------------- comenzi


@restricted
async def cmd_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(HELP + f"\n\n🆔 Chat ID-ul tău: <code>{update.effective_chat.id}</code>")


@restricted
async def cmd_add(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not context.args or not context.args[0].startswith("http"):
        await update.message.reply_text("Folosire: /add &lt;link&gt; [nume]")
        return
    await _add_from_text(update, context, context.args[0], " ".join(context.args[1:]) or None)


@restricted
async def on_text(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """Un link trimis simplu = /add."""
    url = next((w for w in update.message.text.split() if w.startswith("http")), None)
    if url:
        await _add_from_text(update, context, url, None)


async def _add_from_text(update, context, url, name):
    msg = await update.message.reply_text("⏳ Caut prețul...")
    lines = await add_items(context.application, update.effective_chat.id, [{"url": url, "name": name}])
    await msg.edit_text(lines[0] + "\n\nSetează o alertă cu /tinta &lt;id&gt; &lt;preț&gt;")


@restricted
async def cmd_list(update: Update, context: ContextTypes.DEFAULT_TYPE):
    rows = db.list_products(update.effective_chat.id)
    if not rows:
        await update.message.reply_text("Nu urmăresc niciun produs. Trimite-mi un link!")
        return
    lines = []
    for r in rows:
        target = f" 🎯{fmt_price(r['target_price'])}" if r["target_price"] else ""
        warn = " ⚠️" if r["fail_count"] else ""
        lines.append(f"<code>#{r['id']}</code> {link(r)}: <b>{fmt_price(r['last_price'], r['currency'])}</b>{target}{stock_icon(r['in_stock'])}{warn}")
    await send_long(context.bot, update.effective_chat.id, lines, f"📋 <b>{len(rows)} produse</b>\n\n")


@restricted
async def cmd_pret(update: Update, context: ContextTypes.DEFAULT_TYPE):
    rows = db.list_products(update.effective_chat.id)
    if not rows:
        await update.message.reply_text("Nu urmăresc niciun produs.")
        return
    msg = await update.message.reply_text(f"⏳ Verific {len(rows)} produse în paralel...")
    t0 = time.monotonic()
    async with check_lock:  # nu dublăm verificarea programată, dacă rulează chiar acum
        await check_rows(context.application, rows)
    await msg.delete()
    await cmd_list(update, context)
    await update.message.reply_text(f"⏱ {time.monotonic() - t0:.1f}s")


@restricted
async def cmd_remove(update: Update, context: ContextTypes.DEFAULT_TYPE):
    pid = parse_id(context)
    ok = pid is not None and db.remove_product(update.effective_chat.id, pid)
    await update.message.reply_text(f"🗑 Șters #{pid}" if ok else "Folosire: /sterge &lt;id&gt; (vezi /lista)")


@restricted
async def cmd_target(update: Update, context: ContextTypes.DEFAULT_TYPE):
    pid = parse_id(context)
    row = db.get_product(update.effective_chat.id, pid) if pid is not None else None
    try:
        price = float(context.args[1].replace(",", "."))
    except (IndexError, ValueError):
        row = None
    if not row:
        await update.message.reply_text("Folosire: /tinta &lt;id&gt; &lt;preț&gt;  (ex: /tinta 3 450)")
        return
    db.update_fields(pid, target_price=price or None)
    if price:
        await update.message.reply_text(f"🎯 Te anunț când {link(row)} ajunge la ≤ {fmt_price(price, row['currency'])}")
    else:
        await update.message.reply_text(f"Țintă ștearsă pentru {link(row)}")


@restricted
async def cmd_selector(update: Update, context: ContextTypes.DEFAULT_TYPE):
    pid = parse_id(context)
    row = db.get_product(update.effective_chat.id, pid) if pid is not None else None
    if not row or len(context.args) < 2:
        await update.message.reply_text("Folosire: /selector &lt;id&gt; &lt;selector css&gt;")
        return
    db.update_fields(pid, selector=" ".join(context.args[1:]))
    res = await check_product(context.application, db.get_product(update.effective_chat.id, pid), notify=False)
    await update.message.reply_text(
        f"✅ Preț găsit: <b>{fmt_price(res.price, res.currency)}</b> ({res.method})" if res.price else f"⚠️ Tot nu găsesc prețul ({res.error})"
    )


@restricted
async def cmd_history(update: Update, context: ContextTypes.DEFAULT_TYPE):
    pid = parse_id(context)
    row = db.get_product(update.effective_chat.id, pid) if pid is not None else None
    if not row:
        await update.message.reply_text("Folosire: /istoric &lt;id&gt;")
        return
    stats = db.price_stats(pid)
    lines = [
        f"{time.strftime('%d.%m.%Y %H:%M', time.localtime(h['ts']))}  <b>{fmt_price(h['price'], row['currency'])}</b>{stock_icon(h['in_stock'])}"
        for h in db.history(pid)
    ]
    header = (
        f"📊 {link(row)}\nMin: <b>{fmt_price(stats['min'], row['currency'])}</b> · "
        f"Max: {fmt_price(stats['max'], row['currency'])} · {stats['n']} schimbări\n\n"
    )
    await send_long(context.bot, update.effective_chat.id, lines or ["(fără date încă)"], header)


@restricted
async def cmd_status(update: Update, context: ContextTypes.DEFAULT_TYPE):
    rows = db.list_products()
    mine = [r for r in rows if r["chat_id"] == update.effective_chat.id]
    failing = sum(1 for r in mine if r["fail_count"])
    last = max((r["last_checked"] or 0 for r in rows), default=0)
    uptime = int(time.time() - started_at)
    await update.message.reply_text(
        f"✅ Rulez de {uptime // 3600}h {uptime % 3600 // 60}m\n"
        f"📦 Produse: {len(mine)} (cu probleme: {failing})\n"
        f"⏰ Verific la fiecare {config.CHECK_INTERVAL_MIN} min\n"
        f"🕐 Ultima verificare: {time.strftime('%H:%M:%S', time.localtime(last)) if last else '—'}"
    )


def parse_table(filename, data):
    """xlsx/csv cu coloanele nume, url, selector (doar url e obligatoriu)."""
    if filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        rows = load_workbook(io.BytesIO(data), read_only=True).active.iter_rows(values_only=True)
    else:
        rows = csv.reader(io.StringIO(data.decode("utf-8-sig")))
    header = [str(h or "").strip().lower() for h in next(rows, [])]
    items = []
    for values in rows:
        rec = {k: (str(v).strip() if v is not None else "") for k, v in zip(header, values)}
        if rec.get("url", "").startswith("http"):
            items.append({"url": rec["url"], "name": rec.get("nume") or rec.get("name") or None, "selector": rec.get("selector") or None})
    return items


async def _import(update, context, filename, data):
    items = parse_table(filename, data)
    if not items:
        await update.message.reply_text("Nu am găsit rânduri cu o coloană <code>url</code> validă.")
        return
    msg = await update.message.reply_text(f"⏳ Import {len(items)} produse...")
    lines = await add_items(context.application, update.effective_chat.id, items)
    await msg.delete()
    await send_long(context.bot, update.effective_chat.id, lines, f"📥 <b>Importat {len(items)} produse</b>\n\n")


@restricted
async def on_document(update: Update, context: ContextTypes.DEFAULT_TYPE):
    doc = update.message.document
    if not doc.file_name.lower().endswith((".xlsx", ".csv")):
        await update.message.reply_text("Trimite un fișier .xlsx sau .csv")
        return
    data = bytes(await (await doc.get_file()).download_as_bytearray())
    await _import(update, context, doc.file_name, data)


@restricted
async def cmd_import(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """Ex: /import https://docs.google.com/spreadsheets/d/<ID>/export?format=csv"""
    if not context.args:
        await update.message.reply_text("Folosire: /import &lt;link CSV&gt; sau trimite un fișier .xlsx/.csv")
        return
    url = context.args[0]
    if "docs.google.com/spreadsheets" in url and "export" not in url:
        url = url.split("/edit")[0] + "/export?format=csv"
    async with AsyncSession(impersonate="chrome", timeout=config.REQUEST_TIMEOUT) as s:
        resp = await s.get(url, allow_redirects=True)
    if resp.status_code != 200:
        await update.message.reply_text(f"Nu pot descărca tabelul (HTTP {resp.status_code}). E public?")
        return
    await _import(update, context, "sheet.csv", resp.content)


async def on_error(update: object, context: ContextTypes.DEFAULT_TYPE):
    log.error("Eroare neprinsă", exc_info=context.error)


# ---------------------------------------------------------------- pornire / oprire


async def notify_admins(app, text):
    for chat_id in config.ALLOWED_CHAT_IDS:
        try:
            await app.bot.send_message(chat_id, text)
        except Exception as e:
            log.warning("Nu pot trimite către %s: %s", chat_id, e)


async def post_init(app: Application):
    app.bot_data["scraper"] = await Scraper().__aenter__()
    await app.bot.set_my_commands([
        ("lista", "Produsele urmărite"),
        ("pret", "Verifică acum toate prețurile"),
        ("add", "Adaugă produs"),
        ("tinta", "Alertă la preț țintă"),
        ("istoric", "Istoricul prețului"),
        ("sterge", "Șterge produs"),
        ("status", "Starea botului"),
        ("help", "Ajutor"),
    ])
    await notify_admins(app, f"🚀 Monitorul de prețuri a pornit ({len(db.list_products())} produse, la {config.CHECK_INTERVAL_MIN} min).")


async def post_stop(app: Application):
    await notify_admins(app, f"🛑 Monitorul de prețuri s-a oprit la {time.strftime('%Y-%m-%d %H:%M:%S')}.")


async def post_shutdown(app: Application):
    await app.bot_data["scraper"].__aexit__(None, None, None)


def main():
    if not config.TELEGRAM_TOKEN:
        raise SystemExit("Lipsește TELEGRAM_TOKEN. Copiază .env.example în .env și completează-l.")
    app = (
        Application.builder()
        .token(config.TELEGRAM_TOKEN)
        .defaults(Defaults(parse_mode=ParseMode.HTML, link_preview_options=LinkPreviewOptions(is_disabled=True)))
        .post_init(post_init)
        .post_stop(post_stop)
        .post_shutdown(post_shutdown)
        .concurrent_updates(True)
        .build()
    )
    for names, handler in [
        (("start", "help"), cmd_start),
        (("add",), cmd_add),
        (("lista", "list"), cmd_list),
        (("pret", "check"), cmd_pret),
        (("sterge", "remove"), cmd_remove),
        (("tinta", "target"), cmd_target),
        (("selector",), cmd_selector),
        (("istoric", "history"), cmd_history),
        (("status", "test"), cmd_status),
        (("import",), cmd_import),
    ]:
        app.add_handler(CommandHandler(names, handler))
    app.add_handler(MessageHandler(filters.Document.ALL, on_document))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, on_text))
    app.add_error_handler(on_error)

    app.job_queue.run_repeating(check_all_job, interval=config.CHECK_INTERVAL_MIN * 60, first=20)
    log.info("Bot pornit. Interval: %d min", config.CHECK_INTERVAL_MIN)
    app.run_polling(allowed_updates=Update.ALL_TYPES, drop_pending_updates=True)


if __name__ == "__main__":
    main()
